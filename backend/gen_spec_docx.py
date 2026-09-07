# -*- coding: utf-8 -*-
"""Generate the project requirements .docx used as agent task input."""
from docx import Document
from docx.shared import Pt

doc = Document()
doc.add_heading('任务清单管理器 todo-cli 需求规格说明书', 0)

doc.add_heading('1. 项目背景', level=1)
doc.add_paragraph(
    '为个人用户提供一个命令行任务清单管理工具,数据持久化为 JSON 文件,'
    '零第三方运行时依赖(仅 Python 标准库)。本文档是唯一需求来源,'
    '实现必须逐条满足下列功能需求与验收标准。')

doc.add_heading('2. 功能需求', level=1)
frs = [
    ('FR1 命令行接口',
     '程序为 todo.py,通过 python todo.py <子命令> 调用。'
     '子命令包括 add、list、done、delete、stats 五个。'),
    ('FR2 数据存储',
     '数据保存于脚本同目录的 tasks.json;首次运行自动创建;格式为对象数组。'),
    ('FR3 任务字段',
     '每条任务包含:id(自增整数)、title(字符串)、priority(high/medium/low,默认 medium)、'
     'done(布尔值,默认 false)、created_at(ISO8601 字符串)。'),
    ('FR4 add 与 list',
     'add 接收 --title(必填)与 --priority(可选);id 从 1 自增,删除后不复用。'
     'list 输出全部任务,支持 --filter done|undone;排序规则:未完成在前,'
     '同状态按优先级 high>medium>low,再按 created_at 升序。'),
    ('FR5 done 与 delete',
     'done <id> 将任务标记完成;delete <id> 删除任务;对不存在 id 输出错误信息并以非零码退出。'),
    ('FR6 stats',
     'stats 输出三行:total(总数)、done(完成数)、completion_rate(完成率,百分比保留 1 位小数)。'),
]
for title, body in frs:
    doc.add_heading(title, level=2)
    doc.add_paragraph(body)

doc.add_heading('3. 验收标准', level=1)
acs = [
    'A1: 自测脚本 test_todo.py 使用 unittest,至少覆盖 8 个用例,'
    '覆盖 add/list 排序/done/delete/stats/不存在 id 的错误处理。',
    'A2: 运行 python test_todo.py 全部用例通过,退出码 0。',
    'A3: 代码仅使用 Python 标准库;todo.py 具备 main 入口,'
    '可被 python todo.py 直接执行。',
    'A4: README.md 简要说明五个子命令的用法(每个至少一行示例)。',
]
for a in acs:
    doc.add_paragraph(a, style='List Bullet')

doc.add_heading('4. 交付物', level=1)
doc.add_paragraph('todo.py、test_todo.py、README.md,置于 todo_cli/ 目录;另需一份 project_report.pptx 汇报 PPT。')

doc.save('project_spec.docx')
print('saved project_spec.docx')
