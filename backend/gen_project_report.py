#!/usr/bin/env python3
"""Generate project_report.pptx for the todo-cli project."""
from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()

# Slide 1: Title
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "todo-cli 项目报告"
slide.placeholders[1].text = "D:/AI/code/cognitive-os/backend/todo_cli"

# Slide 2: Project background
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "项目背景"
tf = slide.placeholders[1].text_frame
tf.text = (
    "todo-cli 是一个基于 Python 标准库的命令行任务管理器。\n"
    "它提供 add/list/done/delete/stats 五个子命令，\n"
    "数据以 JSON 形式持久化到 tasks.json，无需任何第三方依赖。"
)

# Slide 3: Architecture design
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "架构设计"
tf = slide.placeholders[1].text_frame
tf.text = (
    "- 入口: main() 使用 argparse 解析子命令\n"
    "- 数据层: load_tasks()/save_tasks() 读写 tasks.json\n"
    "- 任务字段: id/title/priority/done/created_at\n"
    "- list 排序: 未完成在前 -> 优先级 high>medium>low -> created_at 升序\n"
    "- 仅使用 Python 标准库"
)

# Slide 4: Five subcommands
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "五个子命令说明"
tf = slide.placeholders[1].text_frame
tf.text = (
    "1. add --title <标题> [--priority high|medium|low]  添加任务\n"
    "2. list [--filter done|undone]  列出任务（按规则排序）\n"
    "3. done <id>  将任务标记为已完成\n"
    "4. delete <id>  删除指定任务\n"
    "5. stats  显示总数/已完成数/完成率"
)

# Slide 5: Test results
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "真实测试通过数"
tf = slide.placeholders[1].text_frame
tf.text = (
    "Ran 11 tests in 0.104s\n"
    "OK\n"
    "EXIT_CODE=0\n"
    "\n"
    "全部 11 个 unittest 用例通过，退出码为 0。"
)

# Slide 6: Acceptance criteria
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "验收标准 A1-A4 达成情况"
tf = slide.placeholders[1].text_frame
tf.text = (
    "A1: 11 个 unittest 用例，覆盖全部功能与错误处理 ✅\n"
    "A2: python test_todo.py 全部通过，退出码 0 ✅\n"
    "A3: 仅标准库，todo.py 可被 python todo.py 直接执行 ✅\n"
    "A4: README.md 已存在 ✅"
)

out = "D:/AI/code/cognitive-os/backend/project_report.pptx"
prs.save(out)
print("saved:", out)
