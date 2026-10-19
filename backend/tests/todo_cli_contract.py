"""Independent acceptance checks for the real todo-cli coding task.

The generated unit tests are useful, but they are not trusted as the sole
source of truth: this verifier exercises the command line against a fresh
copy and checks the requirement that deleted IDs are never reused.
"""
from __future__ import annotations

import argparse
import ast
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any


def extract_pptx_text(presentation: Any) -> tuple[int, str]:
    """Read slide text from shape objects (slides themselves have no .text)."""
    parts: list[str] = []
    slides = presentation.slides
    for slide in slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                parts.append(shape.text)
    return len(slides), "\n".join(parts)


def _run(script: Path, *args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, str(script), *args],
        cwd=script.parent,
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=True,
        timeout=30,
        check=False,
    )


def _regression_test_result(test_file: Path) -> dict[str, Any]:
    """Require a real test for delete-then-add with the exact expected ID 3."""
    if not test_file.is_file():
        return {"passed": False, "error": "missing test_todo.py"}
    try:
        tree = ast.parse(test_file.read_text(encoding="utf-8"))
    except (OSError, SyntaxError) as exc:
        return {"passed": False, "error": f"cannot parse test_todo.py: {exc}"}
    for node in ast.walk(tree):
        if not isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)) or node.name != "test_id_not_reused_after_delete":
            continue
        additions = sum(
            isinstance(child, ast.Call)
            and isinstance(child.func, ast.Attribute)
            and child.func.attr == "add_task"
            for child in ast.walk(node)
        )
        deletes = sum(
            isinstance(child, ast.Call)
            and isinstance(child.func, ast.Attribute)
            and child.func.attr == "delete_task"
            for child in ast.walk(node)
        )
        expects_three = any(
            isinstance(child, ast.Call)
            and isinstance(child.func, ast.Attribute)
            and child.func.attr in ("assertEqual", "assertEquals")
            and len(child.args) >= 2
            and isinstance(child.args[1], ast.Constant)
            and child.args[1].value == 3
            and isinstance(child.args[0], ast.Subscript)
            and isinstance(child.args[0].slice, ast.Constant)
            and child.args[0].slice.value == "id"
            for child in ast.walk(node)
        )
        passed = additions >= 3 and deletes >= 1 and expects_three
        return {
            "passed": passed,
            "test_name": node.name,
            "add_calls": additions,
            "delete_calls": deletes,
            "expects_id_3": expects_three,
            **({} if passed else {"error": "test must add twice, delete, add again, and assert the new task ID is 3"}),
        }
    return {"passed": False, "error": "missing test_id_not_reused_after_delete"}


def verify_todo_project(project_dir: Path, pptx_path: Path | None = None) -> dict[str, Any]:
    """Run self-tests and an independent add/add/delete/add acceptance case."""
    project_dir = project_dir.resolve()
    result: dict[str, Any] = {
        "passed": False,
        "self_tests": {"passed": False},
        "regression_test": {"passed": False},
        "id_not_reused": {"passed": False, "observed_ids": [], "expected_ids": [1, 3]},
    }
    source = project_dir / "todo.py"
    if not source.is_file():
        result["error"] = f"missing {source.name}"
        return result

    tests = project_dir / "test_todo.py"
    result["regression_test"] = _regression_test_result(tests)
    if tests.is_file():
        with tempfile.TemporaryDirectory(prefix="todo-cli-self-test-") as tmp:
            sandbox = Path(tmp)
            shutil.copy2(source, sandbox / "todo.py")
            shutil.copy2(tests, sandbox / "test_todo.py")
            check = _run(sandbox / "test_todo.py")
            result["self_tests"] = {
                "passed": check.returncode == 0,
                "return_code": check.returncode,
                "output": (check.stdout + check.stderr)[-4000:],
            }
    else:
        result["self_tests"] = {"passed": False, "error": "missing test_todo.py"}

    with tempfile.TemporaryDirectory(prefix="todo-cli-contract-") as tmp:
        sandbox = Path(tmp)
        script = sandbox / "todo.py"
        shutil.copy2(source, script)
        calls = [
            _run(script, "add", "--title", "first"),
            _run(script, "add", "--title", "deleted highest id"),
            _run(script, "delete", "2"),
            _run(script, "add", "--title", "after deletion"),
        ]
        failed = next((call for call in calls if call.returncode != 0), None)
        if failed:
            result["id_not_reused"]["error"] = (failed.stdout + failed.stderr)[-2000:]
        else:
            try:
                tasks = json.loads((sandbox / "tasks.json").read_text(encoding="utf-8"))
                observed = [task["id"] for task in tasks]
                result["id_not_reused"]["observed_ids"] = observed
                result["id_not_reused"]["passed"] = observed == [1, 3]
                if observed != [1, 3]:
                    result["id_not_reused"]["error"] = (
                        "add/add/delete(2)/add must leave task IDs [1, 3]; "
                        f"observed {observed}"
                    )
            except (OSError, json.JSONDecodeError, KeyError, TypeError) as exc:
                result["id_not_reused"]["error"] = f"invalid tasks.json after CLI scenario: {exc}"

    if pptx_path is not None:
        try:
            from pptx import Presentation

            slide_count, text = extract_pptx_text(Presentation(str(pptx_path.resolve())))
            result["pptx"] = {"opens": True, "slides": slide_count, "has_text": bool(text.strip())}
        except Exception as exc:  # include the actual diagnostic in evaluation output
            result["pptx"] = {"opens": False, "error": f"{type(exc).__name__}: {exc}"}
    result["passed"] = bool(
        result["self_tests"].get("passed")
        and result["regression_test"].get("passed")
        and result["id_not_reused"].get("passed")
        and (pptx_path is None or result.get("pptx", {}).get("opens"))
    )
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-dir", required=True, type=Path)
    parser.add_argument("--pptx", type=Path)
    args = parser.parse_args()
    result = verify_todo_project(args.project_dir, args.pptx)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
