"""Opt-in real-model long task: source research -> technical PPT -> self-verification.

Uses only the configured provider in a new isolated agent process. The harness
never writes or repairs its presentation. Run with --config <installed config>.
"""
import argparse
import collections
import hashlib
import json
import os
from pathlib import Path
import socket
import subprocess
import tempfile
import time
import urllib.request
import zipfile

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--config', type=Path, required=True)
parser.add_argument('--timeout', type=int, default=1200)
parser.add_argument('--reuse-output', type=Path, help='Continue verification of artifacts from a previous real run; never regenerate them in the harness')
args = parser.parse_args()
backend = Path(__file__).resolve().parents[1]
output = args.reuse_output.resolve() if args.reuse_output else backend/'build'/time.strftime('live-long-task-%Y%m%d-%H%M%S')
output.mkdir(parents=True,exist_ok=bool(args.reuse_output))
if args.reuse_output:
    for name in ('cognitive-os-context-runtime.pptx','technical-qa.md','verification.json'):
        if not (output/name).is_file(): raise SystemExit('Missing prior agent artifact: '+name)
    for name in ('final-task.json','run-summary.json','samples.jsonl','server.log','task.txt','monitor.json','independent-verification.json'):
        if (output/name).exists(): (output/name).rename(output/('baseline-'+name))
original = json.loads(args.config.read_text(encoding='utf-8'))
config = {key:value for key,value in original.items() if key.startswith('llm.')}
if config.get('llm.provider') in (None,'mock') or not config.get('llm.api_key'):
    raise SystemExit('A configured real API is required')
# No production round/request/token cap is introduced by this test.
config['reasoning.max_rounds'] = -1
config['llm.routing'] = 'fixed'
sources = ['include/context/mmu.h','src/context/mmu.c','src/retrieval/context_builder.c',
           'src/action/tool_file.c','src/cognition/planner.c','src/cognition/reasoning.c']
prompt = f'''根据最新 Cognitive OS Agent 源码生成一份 20–24 页中文技术深度 PPT，适合完整技术讲解以及别人提问。
只读源码，不修改项目文件。所有写入、脚本、产物都只能放在当前工作目录 {output.as_posix()}。
需要实际阅读的源码：{json.dumps([(backend/p).as_posix() for p in sources],ensure_ascii=False)}。
mmu.c、context_builder.c、tool_file.c、planner.c 要完成相关完整分页阅读；reasoning.c 可以先搜索并阅读任务循环、失败恢复和完成判定的相关函数，不要求全文。
内容必须依据真实实现，覆盖分层记忆与上下文、页缓存与版本失效、检索组装、工具执行、模型上下文及输出预算、无进展检测、失败恢复和路径重新发现、完成证据、并发与局限。
输出 cognitive-os-context-runtime.pptx（20–24页，每页有具体机制、源码函数或数据流，至少8页标注源码路径/函数；至少18页正文超过60字，至少16页演讲者备注超过100字）、technical-qa.md（至少3000字，包含不少于15个深入技术问题、源码依据、取舍与回答）、verification.json。
可以使用已安装的 python-pptx。自行编写生成脚本，实际执行；脚本失败应根据真实错误修复后重跑。不要只写脚本或只描述计划。
完成前自行用 Python 打开 PPTX，检查页数、文字、演讲备注、源码引用，并将检查结果写入 verification.json。最终给出实际文件位置及验证结果，未完成不要宣布完成。
避免重复读取未变化且已获取的同一页；每一步应基于现有结果推进。'''
if args.reuse_output:
    prompt=('继续完成上一次真实 Agent 任务的最终核验。当前工作目录已有 Agent 自主生成的 PPT、问答和验证记录，'
            '上一轮任务的实际状态、失败原因和执行证据见 baseline-final-task.json；不要自行假设原因。'
            '上次真实执行的统计保存在 baseline-run-summary.json。请复用已有产物及已完成的源码研究，'
            '先实际打开并独立核对 PPT 页数、正文、备注、引用和问答字数；合格的 PPT 不要重做，不要再次逐页读取已经研究完成的源码。'
            '只在需要时阅读相关片段，不要求全文。更新 verification.json，交付实际文件及核验结果。'
            '原任务的产物和内容要求仍然如下（其中源码阅读已在上一轮完成，直接复用）：\n'+prompt)
(output/'task.txt').write_text(prompt,encoding='utf-8')

with tempfile.TemporaryDirectory(prefix='coa-live-long-state-') as directory:
    state = Path(directory)
    (state/'cognitive-os-agent.json').write_text(json.dumps(config),encoding='utf-8')
    with socket.socket() as reservation:
        reservation.bind(('127.0.0.1',0)); port = reservation.getsockname()[1]
    def request(path, body=None):
        req=urllib.request.Request(f'http://127.0.0.1:{port}'+path,
            data=json.dumps(body).encode() if body is not None else None,
            headers={'Content-Type':'application/json'})
        with urllib.request.urlopen(req,timeout=10) as response: return json.load(response)
    binary=backend/'build'/('cognitive-os-agent.exe' if os.name=='nt' else 'cognitive-os-agent')
    seen_ids=set(); signatures=collections.Counter(); failures=[]; task={}
    with (output/'server.log').open('w',encoding='utf-8') as log, (output/'samples.jsonl').open('w',encoding='utf-8') as samples:
        process=subprocess.Popen([str(binary),'--state',str(state),'--workspace',str(output),'serve',str(port)],
            cwd=backend,stdout=log,stderr=log,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
        try:
            for _ in range(100):
                try: request('/v1/scheduler'); break
                except OSError:
                    if process.poll() is not None: raise RuntimeError('Isolated agent exited during startup')
                    time.sleep(.1)
            task_id=request('/v1/tasks',{'prompt':prompt})['id']
            (output/'monitor.json').write_text(json.dumps({'port':port,'task_id':task_id,'process_id':process.pid,'model':config.get('llm.model')}),encoding='utf-8')
            print(json.dumps({'output':str(output),'port':port,'task_id':task_id,'model':config.get('llm.model')},ensure_ascii=False),flush=True)
            started=time.monotonic(); last_print=-30
            while time.monotonic()-started < args.timeout:
                task=request('/v1/tasks/'+str(task_id))
                for step in task.get('steps',[]):
                    if step.get('id') in seen_ids: continue
                    seen_ids.add(step.get('id'))
                    evidence=step.get('out') or ''
                    if '\n' in evidence and '[bytes ' in evidence:
                        evidence=evidence.split('\n',1)[1]
                    signature=(step.get('tool'),step.get('args'),hashlib.sha256(evidence.encode()).hexdigest())
                    signatures[signature]+=1
                    if step.get('ok')==0: failures.append({key:step.get(key) for key in ('id','tool','target')})
                summary={key:task.get(key) for key in ('status','stage','round','llm_calls','tool_calls','model_failures','tokens_in','tokens_out','usage_available','elapsed_ms')}
                summary.update(unique_observations=len(signatures),repeated_observations=sum(n-1 for n in signatures.values()),failed_actions=len(failures))
                samples.write(json.dumps(summary,ensure_ascii=False)+'\n'); samples.flush()
                if time.monotonic()-started-last_print >= 30:
                    print(json.dumps(summary,ensure_ascii=False),flush=True); last_print=time.monotonic()-started
                if task.get('status') not in ('RUNNING','QUEUED'): break
                time.sleep(2)
            (output/'final-task.json').write_text(json.dumps(task,ensure_ascii=False,indent=2),encoding='utf-8')
            (output/'run-summary.json').write_text(json.dumps({**summary,'model':config.get('llm.model'),'failed_steps':failures},ensure_ascii=False,indent=2),encoding='utf-8')
            assert task.get('status')=='DONE', 'Long task did not complete; inspect saved task and samples'
            deck=output/'cognitive-os-context-runtime.pptx'
            with zipfile.ZipFile(deck) as archive: assert archive.testzip() is None
            from pptx import Presentation
            presentation=Presentation(deck)
            texts=['\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame) for slide in presentation.slides]
            notes=[slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else '' for slide in presentation.slides]
            qa=(output/'technical-qa.md').read_text(encoding='utf-8')
            checks={'slides':len(texts),'substantive_slides':sum(len(text)>=60 for text in texts),
                'notes_over_100_chars':sum(len(note)>=100 for note in notes),
                'cited_slides':sum(any(token in text for token in ('src/','include/','mmu.c','reasoning.c','planner.c','tool_file.c')) for text in texts),
                'qa_chars':len(qa),'verification_json_valid':isinstance(json.loads((output/'verification.json').read_text(encoding='utf-8')),dict)}
            (output/'independent-verification.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
            print(json.dumps(checks,ensure_ascii=False),flush=True)
            assert 20<=checks['slides']<=24 and checks['substantive_slides']>=18
            assert checks['notes_over_100_chars']>=16 and checks['cited_slides']>=8 and checks['qa_chars']>=3000
            print('PASS: real model completed and independently verified the long technical presentation task.',flush=True)
        finally:
            process.terminate()
            try: process.wait(timeout=5)
            except subprocess.TimeoutExpired: process.kill(); process.wait()
