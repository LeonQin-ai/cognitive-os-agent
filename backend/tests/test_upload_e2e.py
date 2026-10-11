"""Real HTTP + scheduler + LLM adapter + persistence + browser regression run.
Uses a local capture provider: no real API keys or paid model requests.
Run after build.ps1 -Target cli and build.ps1 -Target review.
"""
import base64
import io
import json
import os
from pathlib import Path
import socket
import subprocess
import tempfile
import threading
import time
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

BACKEND = Path(__file__).resolve().parents[1]
captures = []
checks = []
task_ids = []


class Provider(BaseHTTPRequestHandler):
    def log_message(self, *args): pass

    def do_POST(self):
        body = json.loads(self.rfile.read(int(self.headers['Content-Length'])))
        captures.append(body)
        text = '附件内容已核验，测试回答完成。'
        if self.path.endswith('/messages'):
            reply = {'id': 'test', 'type': 'message', 'role': 'assistant',
                     'content': [{'type': 'text', 'text': text}], 'stop_reason': 'end_turn',
                     'usage': {'input_tokens': 10, 'output_tokens': 10}}
        else:
            reply = {'choices': [{'message': {'role': 'assistant', 'content': text},
                                  'finish_reason': 'stop'}],
                     'usage': {'prompt_tokens': 10, 'completion_tokens': 10}}
        data = json.dumps(reply, ensure_ascii=False).encode()
        self.send_response(200); self.send_header('Content-Type', 'application/json')
        self.send_header('Content-Length', str(len(data))); self.end_headers(); self.wfile.write(data)


def check(label, condition):
    assert condition, label
    checks.append(label)
    print('PASS', label, flush=True)


def request(method, path, body=None, status=200):
    if isinstance(body, dict): body = json.dumps(body).encode()
    req = urllib.request.Request(base + path, data=body, method=method)
    try:
        with urllib.request.urlopen(req, timeout=90) as response:
            actual, data = response.status, response.read()
    except urllib.error.HTTPError as response:
        actual, data = response.code, response.read()
    assert actual == status, (path, actual, data[:500])
    return json.loads(data) if data else None


def upload(name, data, session='A', status=200):
    query = urllib.parse.urlencode({'name': name, 'session': session})
    return request('POST', '/v1/upload?' + query, data, status)


def chat(ids, session='A', expected=200):
    start = len(captures)
    response = request('POST', '/v1/chat', {'message': '请描述并总结附件内容',
                                          'session': session, 'attachments': ids}, expected)
    if expected != 200: return
    task_ids.append(response['id'])
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline:
        task = request('GET', '/v1/tasks/' + str(response['id']))
        if task['status'] in ('DONE', 'FAILED', 'CANCELLED', 'TIMEOUT'):
            assert task['status'] == 'DONE', task
            check('chat task completed with captured provider request', len(captures) > start)
            return captures[start:]
        time.sleep(.05)
    raise AssertionError('chat task timed out')


def office(name, xml):
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w') as archive: archive.writestr(name, xml)
    return out.getvalue()


def pdf(text):
    stream = f'BT /F1 12 Tf 10 100 Td ({text}) Tj ET'.encode()
    objects = [b'<< /Type /Catalog /Pages 2 0 R >>', b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
               b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
               b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
               b'<< /Length ' + str(len(stream)).encode() + b' >>\nstream\n' + stream + b'\nendstream']
    data = b'%PDF-1.4\n'; offsets = []
    for i, obj in enumerate(objects, 1):
        offsets.append(len(data)); data += str(i).encode() + b' 0 obj\n' + obj + b'\nendobj\n'
    start = len(data); data += b'xref\n0 6\n0000000000 65535 f \n'
    for offset in offsets: data += f'{offset:010d} 00000 n \n'.encode()
    return data + f'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{start}\n%%EOF'.encode()


def raw(data):
    with socket.create_connection(('127.0.0.1', port), timeout=3) as c:
        c.sendall(data); response = b''
        while True:
            part = c.recv(65536)
            if not part: break
            response += part
    return int(response.split(b' ')[1]), response


provider = ThreadingHTTPServer(('127.0.0.1', 0), Provider)
threading.Thread(target=provider.serve_forever, daemon=True).start()
with socket.socket() as s: s.bind(('127.0.0.1', 0)); port = s.getsockname()[1]
base = f'http://127.0.0.1:{port}'
process = None
with tempfile.TemporaryDirectory(prefix='attachment-e2e-', dir=BACKEND/'build') as tmp:
    state = Path(tmp)
    config = {'llm.provider': 'openai', 'llm.model': 'capture-vision', 'llm.api_key': 'local-test-only',
              'llm.base_url': f'http://127.0.0.1:{provider.server_port}', 'scheduler.workers': 2, 'scheduler.history.count': 4}
    (state/'cognitive-os-agent.json').write_text(json.dumps(config))
    (state/'uploads').mkdir()
    (state/'uploads/legacy.txt').write_text('LEGACY_TEXT_PRESERVED')
    log = open(BACKEND/'build/review-e2e-server.log', 'w', encoding='utf-8')

    def start():
        global process
        process = subprocess.Popen([str(BACKEND/'build/cognitive-os-agent.exe'), '--state', str(state),
                                    '--workspace', str(state), 'serve', str(port)],
                                   stdout=log, stderr=log, creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
        for _ in range(150):
            if process.poll() is not None: raise AssertionError('runtime exited during startup')
            try: request('GET', '/v1/scheduler'); return
            except (OSError, TimeoutError): time.sleep(.1)
        raise AssertionError('runtime not ready')

    def stop():
        process.terminate(); process.wait(timeout=10)

    try:
        start()
        legacy=next(x for x in request('GET','/v1/uploads')['files'] if x['name']=='legacy.txt')
        check('pre-ID attachments migrate without losing bytes',legacy['status']=='ready' and (state/'uploads'/f"{legacy['id']}.raw").read_bytes()==b'LEGACY_TEXT_PRESERVED')
        # Real concurrent TCP connections: slow headers cannot block normal requests.
        slow = socket.create_connection(('127.0.0.1', port)); slow.sendall(b'POST /v1/upload HTTP/1.1\r\n')
        before = time.monotonic(); request('GET', '/v1/scheduler')
        check('slow connection does not block API', time.monotonic()-before < 2); slow.close()
        status, body = raw(b'POST /v1/upload?name=lower.txt HTTP/1.1\r\nHost: localhost\r\ncontent-length: 4\r\n\r\ntext')
        check('lowercase Content-Length uploads complete body', status == 200 and b'"size":4' in body)
        for headers in [b'Content-Length: 1\r\nContent-Length: 2\r\n', b'Content-Length: 2x\r\n',
                        b'Transfer-Encoding: chunked\r\n', b'Content-Length: -1\r\n']:
            check('invalid HTTP framing rejected: '+headers.decode().strip(),
                  raw(b'POST /v1/upload?name=a.txt HTTP/1.1\r\nHost: x\r\n'+headers+b'\r\nxx')[0] == 400)
        check('oversized HTTP body rejected before allocation',
              raw(b'POST /v1/upload HTTP/1.1\r\nContent-Length: 99999999999999999999\r\n\r\n')[0] == 413)
        first = upload('same.txt', b'first immutable content'); second = upload('same.txt', b'second content')
        check('same-name files get distinct immutable IDs', first['id'] != second['id'])
        check('original raw file preserved', (state/'uploads'/f"{first['id']}.raw").read_bytes() == b'first immutable content')
        a = upload('a'*140+'.txt', b'x'*1800); b = upload('a'*140+'b.txt', b'y'*1800)
        check('long filenames keep distinct IDs and bounded chunks', a['id'] != b['id'] and a['chunks'] == 3 and b['chunks'] == 3)
        upload('big.txt', b'x'*(64*1024*1024+1), status=413); check('uniform 64 MiB upload limit', True)
        image_bytes = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aK1sAAAAASUVORK5CYII=')
        image = upload('image.png', image_bytes)
        check('PNG saved as ready image, not indexed text', image['status']=='ready' and image['mime']=='image/png' and image['chunks']==0)
        requests = chat([image['id']])
        serialized = json.dumps(requests)
        check('OpenAI actual payload contains exact image bytes', 'image_url' in serialized and base64.b64encode(image_bytes).decode() in serialized)
        chat([image['id']], session='B', expected=400); check('cross-session attachment denied', True)
        chat(['../image'], expected=400); check('invalid attachment ID denied', True)
        docs = [('notes.docx', office('word/document.xml', '<d><t>DOCX_UNIQUE_TEXT</t></d>'), 'DOCX_UNIQUE_TEXT'),
                ('sheet.xlsx', office('xl/sharedStrings.xml', '<d><t>XLSX_UNIQUE_TEXT</t></d>'), 'XLSX_UNIQUE_TEXT'),
                ('slides.pptx', office('ppt/slides/slide1.xml', '<d><t>PPTX_UNIQUE_TEXT</t></d>'), 'PPTX_UNIQUE_TEXT'),
                ('notes.pdf', pdf('PDF_UNIQUE_TEXT'), 'PDF_UNIQUE_TEXT')]
        for name, data, marker in docs:
            obj = upload(name, data)
            if obj['status']!='ready': print('Processing result:',obj,flush=True)
            check(name+' extraction ready', obj['status']=='ready' and obj['chunks']>0)
            requests = chat([obj['id']])
            check(name+' extracted text reached model', marker in json.dumps(requests))
        # Full Office files and native legacy formats, not only XML fixtures.
        from docx import Document
        from pptx import Presentation
        from pptx.util import Inches
        from PIL import Image, ImageDraw
        import pymupdf as fitz
        import cv2
        import numpy as np
        fixture_dir=state/'fixtures'; fixture_dir.mkdir()
        picture=Image.new('RGB',(640,360),'white')
        ImageDraw.Draw(picture).text((40,80),'SCAN_VISUAL_ONLY - RED RECTANGLE',fill='black')
        ImageDraw.Draw(picture).rectangle((200,150,450,300),fill='red')
        png=fixture_dir/'scan.png'; picture.save(png)
        word=Document(); word.add_paragraph('REAL_WORD_CONTENT'); word.add_picture(str(png))
        docx_path=fixture_dir/'real.docx'; word.save(docx_path)
        ppt=Presentation()
        slide=ppt.slides.add_slide(ppt.slide_layouts[5]); slide.shapes.title.text='REAL_PPT_CONTENT'
        slide.shapes.add_picture(str(png),Inches(1),Inches(2),width=Inches(6))
        pptx_path=fixture_dir/'real.pptx'; ppt.save(pptx_path)
        scanner=fitz.open(); page=scanner.new_page(width=640,height=360)
        page.insert_image(page.rect,filename=str(png)); scanned_path=fixture_dir/'scanned.pdf'; scanner.save(scanned_path); scanner.close()
        for name,path in [('real.docx',docx_path),('real.pptx',pptx_path),('scanned.pdf',scanned_path)]:
            if name=='real.docx':
                pending=[]
                worker=threading.Thread(target=lambda: pending.append(upload(name,path.read_bytes())))
                worker.start(); time.sleep(.15)
                before=time.monotonic(); request('GET','/v1/scheduler')
                check('Office processing does not block task queries',time.monotonic()-before<2)
                worker.join(timeout=90); assert pending
                obj=pending[0]
            else: obj=upload(name,path.read_bytes())
            check(name+' produces visual pages' ,obj['status']=='ready' and len(obj.get('frames',[]))>0 and obj.get('page_count',0)>0)
            payload=chat([obj['id']])
            check(name+' visual pages reach model', 'image_url' in json.dumps(payload))
        if os.name=='nt':
            import win32com.client
            app=win32com.client.DispatchEx('Word.Application'); app.Visible=False; app.DisplayAlerts=0; app.AutomationSecurity=3
            doc=None
            try:
                doc=app.Documents.Open(str(docx_path.resolve()),ReadOnly=True,AddToRecentFiles=False)
                legacy_doc=fixture_dir/'real.doc'; doc.SaveAs2(str(legacy_doc.resolve()),FileFormat=0)
            finally:
                if doc is not None: doc.Close(False)
                app.Quit()
            app=win32com.client.DispatchEx('PowerPoint.Application'); app.AutomationSecurity=3; app.DisplayAlerts=1
            deck=None
            try:
                deck=app.Presentations.Open(str(pptx_path.resolve()),ReadOnly=True,Untitled=False,WithWindow=False)
                legacy_ppt=fixture_dir/'real.ppt'; deck.SaveAs(str(legacy_ppt.resolve()),1)
            finally:
                if deck is not None: deck.Close()
                app.Quit()
            for name,path,marker in [('real.doc',legacy_doc,'REAL_WORD_CONTENT'),('real.ppt',legacy_ppt,'REAL_PPT_CONTENT')]:
                obj=upload(name,path.read_bytes())
                check(name+' legacy conversion ready',obj['status']=='ready' and len(obj.get('frames',[]))>0 and obj.get('page_count',0)>0)
                payload=chat([obj['id']])
                check(name+' text and visual content reach model',marker in json.dumps(payload) and 'image_url' in json.dumps(payload))
        bmp=io.BytesIO(); picture.save(bmp,format='BMP')
        obj=upload('picture.bmp',bmp.getvalue())
        check('BMP is normalized for vision',obj['status']=='ready' and len(obj.get('frames',[]))==1)
        check('BMP normalized image reaches model','image_url' in json.dumps(chat([obj['id']])))
        tiff=io.BytesIO(); picture.save(tiff,format='TIFF')
        obj=upload('picture.tiff',tiff.getvalue())
        check('TIFF is normalized for vision',obj['status']=='ready' and len(obj.get('frames',[]))==1)
        import pillow_heif
        pillow_heif.register_heif_opener()
        heic=io.BytesIO(); picture.save(heic,format='HEIF')
        obj=upload('phone.heic',heic.getvalue())
        check('HEIC phone image normalized for vision',obj['status']=='ready' and len(obj.get('frames',[]))==1)
        check('HEIC visual content reaches model','image_url' in json.dumps(chat([obj['id']])))
        obj=upload('vector.svg',b'<svg xmlns="http://www.w3.org/2000/svg" width="320" height="180"><rect width="320" height="180" fill="blue"/></svg>')
        check('SVG normalized for vision',obj['status']=='ready' and len(obj.get('frames',[]))==1)
        gif=io.BytesIO(); picture.save(gif,format='GIF',save_all=True,append_images=[Image.new('RGB',picture.size,'blue')],duration=500,loop=0)
        obj=upload('animated.gif',gif.getvalue())
        check('animated GIF sampled as multiple images',obj['status']=='ready' and len(obj.get('frames',[]))==2)
        check('animated GIF frame labels reach model','Image frame 2' in json.dumps(chat([obj['id']])))
        avi_path=fixture_dir/'clip.avi'
        writer=cv2.VideoWriter(str(avi_path),cv2.VideoWriter_fourcc(*'MJPG'),10,(640,360))
        assert writer.isOpened()
        for i in range(20):
            frame=np.zeros((360,640,3),dtype=np.uint8); frame[:,:,i%3]=180
            cv2.putText(frame,f'FRAME {i} TIME {i/10:.1f}s',(30,100),cv2.FONT_HERSHEY_SIMPLEX,1,(255,255,255),2)
            writer.write(frame)
        writer.release()
        import importlib.util
        spec=importlib.util.spec_from_file_location('coa_processor',BACKEND/'tools/attachment_extract.py')
        processor=importlib.util.module_from_spec(spec); spec.loader.exec_module(processor)
        decoded={'frames':[],'warnings':[]}
        processor.video_fallback(avi_path,str(fixture_dir/'ffmpeg-check'),decoded)
        check('FFmpeg fallback decodes six timestamped images',len(decoded['frames'])==6 and decoded['frames'][-1]['time_ms']>0)
        import imageio_ffmpeg
        ffmpeg=imageio_ffmpeg.get_ffmpeg_exe()
        videos=[avi_path]
        for ext,codec in [('mp4','libx264'),('mov','libx264'),('mkv','libx264'),('webm','libvpx-vp9')]:
            output=fixture_dir/('clip.'+ext)
            subprocess.run([ffmpeg,'-y','-loglevel','error','-i',str(avi_path),'-c:v',codec,'-pix_fmt','yuv420p',str(output)],
                           check=True,capture_output=True,timeout=30)
            videos.append(output)
        video_ids=[]
        video_mimes=[]
        for path in videos:
            obj=upload(path.name,path.read_bytes()); video_ids.append(obj['id']); video_mimes.append(obj['mime'])
            check(path.suffix+' video produces timestamped samples',obj['status']=='ready' and obj['kind']=='video' and
                  len(obj.get('frames',[]))==6 and obj['frames'][-1]['time_ms']>0)
            payload=chat([obj['id']]); data=json.dumps(payload)
            check(path.suffix+' sampled frames and labels reach model','image_url' in data and 'Video frame at' in data and 'Audio is not transcribed' in data)
            urls=[part['image_url']['url'] for req in payload for msg in req['messages']
                  if isinstance(msg.get('content'),list) for part in msg['content'] if part.get('type')=='image_url']
            expected=['data:image/jpeg;base64,'+base64.b64encode((state/'uploads'/frame['file']).read_bytes()).decode()
                      for frame in obj['frames']]
            check(path.suffix+' model image bytes match decoded frames',urls==expected)
        check('video MIME identifies original containers',all(m.startswith('video/') for m in video_mimes))
        combined=chat(video_ids[-4:])
        urls=[part for req in combined for msg in req['messages'] if isinstance(msg.get('content'),list)
              for part in msg['content'] if part.get('type')=='image_url']
        check('four videos deliver all 24 sampled images',len(urls)==24)
        request('DELETE','/v1/uploads/'+video_ids[0])
        check('video deletion cleans generated frame files',not list((state/'uploads').glob(video_ids[0]+'.frame*.jpg')))
        invalid_video=upload('damaged.mp4',b'not a video')
        check('corrupt video is failed instead of fake ready',invalid_video['status']=='failed')
        external_svg=upload('external.svg',b'<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10"><image href="https://example.invalid/image.png"/></svg>')
        check('SVG external resources are rejected',external_svg['status']=='failed')
        invalid = upload('damaged.docx', b'not a zip')
        check('corrupt document reports failed, not ready', invalid['status']=='failed')
        unsupported = upload('binary.bin', b'hello\0binary')
        check('unsupported binary reports unsupported', unsupported['status']=='unsupported')
        chat([invalid['id']], expected=400); check('failed document cannot be submitted', True)
        request('DELETE', '/v1/uploads/'+first['id'])
        check('deleted attachment raw and manifest removed', not (state/'uploads'/f"{first['id']}.raw").exists() and
              not any(x['id']==first['id'] for x in request('GET','/v1/uploads')['files']))
        chat([first['id']], expected=400); check('deleted attachment cannot be reused', True)
        # Runtime history eviction still leaves durable journal records.
        stats=request('GET','/v1/scheduler')
        check('completed task history is bounded in live runtime',stats['retained']<=4)
        request('GET','/v1/tasks/'+str(task_ids[0]),status=404)
        journal=request('GET','/v1/tasks/journal')
        check('evicted tasks remain in durable journal',any(x.get('id')==task_ids[0] for x in journal))
        # Short-lived WS clients stress the HTTP-to-WS socket ownership handoff.
        for _ in range(30):
            with socket.create_connection(('127.0.0.1',port),timeout=3) as ws:
                ws.sendall(b'GET /ws HTTP/1.1\r\nHost: localhost\r\nUpgrade: websocket\r\nConnection: Upgrade\r\nSec-WebSocket-Key: dGhlIHNhbXBsZSBub25jZQ==\r\nSec-WebSocket-Version: 13\r\n\r\n')
                assert b'101 Switching Protocols' in ws.recv(4096)
        request('GET','/v1/scheduler')
        check('short-lived WebSocket ownership handoff preserves HTTP service',True)
        # Browser exercises actual page layout, mocks API timing deterministically.
        env = dict(os.environ, UI_TEST_URL=base, NODE_PATH=str(BACKEND/'webarena/node_modules'), UI_TEST_BROWSER='msedge')
        for script in ['test_progress_ui.cjs', 'test_upload_ui.cjs']:
            subprocess.run(['node', str(BACKEND/'tests'/script)], env=env, check=True, timeout=60)
            check('browser '+script, True)
        before = {x['id']: (x['status'],x['chunks']) for x in request('GET','/v1/uploads')['files']}
        stop(); config['llm.provider']='anthropic'; (state/'cognitive-os-agent.json').write_text(json.dumps(config)); start()
        after = {x['id']: (x['status'],x['chunks']) for x in request('GET','/v1/uploads')['files']}
        check('restart preserves attachment statuses and chunks', before==after)
        check('restored video frame files remain usable','base64' in json.dumps(chat([video_ids[-1]])))
        requests=chat([image['id']]); serialized=json.dumps(requests)
        check('Anthropic actual payload contains exact image bytes', 'base64' in serialized and base64.b64encode(image_bytes).decode() in serialized)
        print(f'E2E: {len(checks)} checks passed', flush=True)
        (BACKEND/'build/review-e2e.json').write_text(json.dumps({'passed':len(checks),'checks':checks},indent=2))
    except Exception:
        print("Runtime exit status:", process.poll(), flush=True)
        raise
    finally:
        if process and process.poll() is None: stop()
        log.close(); provider.shutdown(); provider.server_close()
