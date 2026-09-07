#!/usr/bin/env python3
"""gen_gaia_fixtures.py — generate GAIA-style fixtures (docx + png) for bench_gaia.

Usage: python tools/gen_gaia_fixtures.py <target-dir>
Creates: contact.docx, report.docx (python-docx), diagram.png (Pillow).
Each artifact is optional: if its library is missing, the file is skipped and
bench_gaia marks the corresponding task SKIP instead of FAIL.
"""
import os
import sys


def make_docx(path, heading, paras):
    from docx import Document
    d = Document()
    d.add_heading(heading, level=1)
    for p in paras:
        d.add_paragraph(p)
    d.save(path)
    print(os.path.basename(path), "OK")


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else "gaia"
    os.makedirs(out, exist_ok=True)

    try:
        make_docx(
            os.path.join(out, "contact.docx"),
            "Project Support Guide",
            [
                "For any issue with the deployment, contact the on-call engineer first.",
                "Email: support@example.com",
                "Phone hours: weekdays 9am-6pm.",
            ],
        )
    except Exception as e:
        print("contact.docx SKIP:", e)

    try:
        make_docx(
            os.path.join(out, "report.docx"),
            "Quarterly Revenue Report",
            [
                "Q1 revenue was 120.50 thousand.",
                "Q2 revenue was 155.00 thousand.",
                "Q3 revenue was 169.50 thousand.",
            ],
        )
    except Exception as e:
        print("report.docx SKIP:", e)

    try:
        from PIL import Image, ImageDraw, ImageFont

        img = Image.new("RGB", (480, 160), "white")
        dr = ImageDraw.Draw(img)
        msg = "ACCESS CODE: 7392"
        font = None
        for cand in ("arial.ttf", "DejaVuSans-Bold.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"):
            try:
                font = ImageFont.truetype(cand, 44)
                break
            except Exception:
                continue
        if font is None:
            font = ImageFont.load_default()
        dr.text((20, 55), msg, fill="black", font=font)
        img.save(os.path.join(out, "diagram.png"))
        print("diagram.png OK")
    except Exception as e:
        print("diagram.png SKIP:", e)


if __name__ == "__main__":
    main()
